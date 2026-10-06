"""Сводные таблицы кластеров (этапы 5–6 доработки).

Таблица 1 «Характерные признаки кластеров»: для каждого кластера — средний z-score признаков по его МО,
автоматическое текстовое описание по правилу (пороги — configs/summary.yaml), число МО и примеры.
Тексты, исправленные аналитиками, хранятся в data/cluster_descriptions.yaml вместе с «отпечатком» состава
кластера: если состав изменился, интерфейс предупреждает, что текст писался для другого состава.
"""

from __future__ import annotations

import hashlib
import io
import re

import numpy as np
import pandas as pd
import yaml

from src import clustering, network
from src.io import PROCESSED, ROOT, load_yaml

# палитра кластеров — та же, что в интерфейсе (руководство dataviz): K1…K8, дальше нейтральный серый
CLUSTER_COLORS = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7", "#e34948"]
OTHER_GRAY = "#9a9994"


def cfg() -> dict:
    return load_yaml("summary.yaml")


def cluster_color(label: int) -> str:
    return CLUSTER_COLORS[label] if 0 <= label < len(CLUSTER_COLORS) else OTHER_GRAY


# ============================================================================ разбиение
def partition(netp: dict, method: str, k: int, mode: str = "pooled", seed: int | None = None) -> pd.DataFrame:
    """territory_id, year, label (0 → K1 — самый высокий ВМП на душу в ценах базового года).

    pooled — одна модель на все МО-годы окна (только методы по атрибутам); per_year — каждый год отдельно
    (сеть года; номера K упорядочены по ВМП внутри каждого года).
    """
    cc = clustering.default_params()
    mp = {**cc["methods"][method], "method": method, "k": k, "seed": cc["seed"] if seed is None else seed}
    if mode == "pooled":
        X = network.usable_rows(network.features(netp), netp.get("structural_missing", 0.5))
        lab = clustering.order_labels(clustering.fit_pooled(X, mp), netp["prices"])
        return lab.rename("label").reset_index()
    out = []
    for y, net in network.build(netp).items():
        lab = clustering.fit(net.X.to_numpy(), net.W * net.A, mp)
        lab = clustering.order_labels(pd.Series(lab, index=net.ids), netp["prices"], y)
        out.append(pd.DataFrame({"territory_id": net.ids, "year": y, "label": lab.to_numpy()}))
    return pd.concat(out, ignore_index=True)


def fingerprint(ids) -> str:
    """Отпечаток состава кластера: хэш отсортированного списка territory_id."""
    s = ",".join(str(int(i)) for i in sorted(ids))
    return hashlib.sha1(s.encode()).hexdigest()[:12]


# ============================================================================ таблица 1
def feature_list(netp: dict, c: dict | None = None) -> list[str]:
    c = c or cfg()["characteristic"]
    reg = pd.read_parquet(PROCESSED / "indicator_registry.parquet")["code"]
    out = list(dict.fromkeys([*netp["features"], *c.get("extra_features", [])]))
    return [f for f in out if f in set(reg)]


def zscores(netp: dict, feats: list[str], scope: str = "year") -> pd.DataFrame:
    """z-score признаков (после пересчёта в цены базового года и логарифма по реестру) по МО выборки.

    scope = year — по всем МО выборки отдельно в каждом году; panel — по всем МО-годам окна.
    """
    p = network.merge_params(netp, {"preprocess": {"method": "zscore", "scope": scope, "winsor": None}})
    p["features"] = {f: 1.0 for f in feats}
    return network.features(p)


def real_values(netp: dict, feats: list[str]) -> pd.DataFrame:
    """Исходные значения признаков выборки (денежные — в ценах из netp['prices']), индекс (territory_id, year)."""
    from src import prices

    rows = network.sample_rows(netp)
    pr = netp.get("prices") or {}
    if pr.get("values", "real") == "real":
        rows = prices.to_real(
            rows, feats, pr["base_year"], pr["deflator_scope"], bool(pr.get("spatial_price_adjustment"))
        )
    return rows.set_index(["territory_id", "year"])[feats]


def short_name(name: str) -> str:
    """«ВМП на душу населения (оценка команды)» → «ВМП на душу населения»; первая буква — строчная, кроме аббревиатур."""
    s = re.sub(r"\s*\([^)]*\)\s*$", "", str(name)).strip()
    if len(s) > 1 and not s[1].isupper():
        s = s[0].lower() + s[1:]
    return s


def category(zm: pd.DataFrame, c: int, f: str, th: dict) -> str:
    """max | min | high | low | mid для кластера c и признака f по матрице средних z (кластеры × признаки)."""
    z = zm.at[c, f]
    if not np.isfinite(z):
        return "na"
    if z >= th["extreme"] and zm[f].idxmax() == c:
        return "max"
    if z <= -th["extreme"] and zm[f].idxmin() == c:
        return "min"
    if z >= th["high"]:
        return "high"
    if z <= -th["high"]:
        return "low"
    return "mid"


PHRASES = {
    "max": "Максимальные значения",
    "high": "Высокие значения",
    "mid": "Близкие к среднему",
    "low": "Низкие значения",
    "min": "Минимальные значения",
}


def describe(zm: pd.DataFrame, c: int, names: dict[str, str], c_cfg: dict | None = None) -> tuple[str, list[str]]:
    """Автоматический текст для кластера c и выбранные признаки (3–5 с наибольшим |z|)."""
    c_cfg = c_cfg or cfg()["characteristic"]
    th = c_cfg["thresholds"]
    lo, hi = c_cfg["n_features"]
    z = zm.loc[c].dropna()
    order = z.abs().sort_values(ascending=False).index.tolist()
    strong = [f for f in order if abs(z[f]) >= th["high"]]
    sel = strong[:hi] if len(strong) >= lo else order[: max(lo, len(strong))]
    groups: dict[str, list[str]] = {}
    for f in sel:
        groups.setdefault(category(zm, c, f, th), []).append(names.get(f, f))
    parts = [f"{PHRASES[g]}: {', '.join(groups[g])}." for g in PHRASES if g in groups]
    return " ".join(parts), sel


def characteristic_table(labels: pd.DataFrame, netp: dict, year: int | None = None, c_cfg: dict | None = None) -> dict:
    """Таблица 1 по разбиению labels (territory_id, year, label) за год year (по умолчанию — последний).

    Возвращает {"table": кластер | число МО | текст | примеры | состав, "z": средние z (кластер × признак),
    "raw": средние исходные значения (денежные — в ценах базового года), "features": признаки, "year": год}.
    """
    c_cfg = c_cfg or cfg()["characteristic"]
    year = int(year or labels["year"].max())
    feats = feature_list(netp, c_cfg)
    lab = labels[labels["year"].eq(year)].set_index("territory_id")["label"].astype(int)
    Z = zscores(netp, feats, c_cfg.get("z_scope", "year")).xs(year, level="year").reindex(lab.index)
    R = real_values(netp, feats).xs(year, level="year").reindex(lab.index)
    zm = Z.groupby(lab).mean()
    rm = R.groupby(lab).mean()
    reg = pd.read_parquet(PROCESSED / "indicator_registry.parquet").set_index("code")
    names = {f: short_name(reg.at[f, "name"]) for f in feats}
    mo = pd.read_parquet(PROCESSED / "mo.parquet").set_index("territory_id")
    pop = pd.read_parquet(PROCESSED / "indicators_wide.parquet", columns=["territory_id", "year", "pop"])
    pop = pop[pop["year"].eq(year)].set_index("territory_id")["pop"]
    rows = []
    for c in sorted(lab.unique()):
        ids = lab.index[lab.eq(c)]
        text, sel = describe(zm, c, names, c_cfg)
        ex = pop.reindex(ids).sort_values(ascending=False).head(int(c_cfg.get("n_examples", 4))).index
        rows.append(
            {
                "label": int(c),
                "Кластер": clustering.code(c),
                "Число МО": len(ids),
                "Характерные признаки": text,
                "Примеры МО": ", ".join(mo.loc[ex, "name_short"].astype(str)),
                "признаки": sel,
                "состав": fingerprint(ids),
            }
        )
    return {"table": pd.DataFrame(rows), "z": zm, "raw": rm, "features": feats, "names": names, "year": year}


# ============================================================================ тексты аналитиков
def _desc_path():
    return ROOT / cfg()["descriptions_file"]


def load_descriptions() -> dict:
    p = _desc_path()
    return (yaml.safe_load(p.read_text(encoding="utf-8")) or {}) if p.exists() else {}


def save_description(key: str, code: str, text: str, members: str, auto: str) -> None:
    """Сохранить исправленный текст кластера code для разбиения key (вместе с отпечатком состава)."""
    allx = load_descriptions()
    allx.setdefault(key, {})[code] = {"text": text, "состав": members, "авто": auto}
    p = _desc_path()
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(yaml.safe_dump(allx, allow_unicode=True, sort_keys=True), encoding="utf-8")


def apply_descriptions(table: pd.DataFrame, key: str) -> pd.DataFrame:
    """Подставить тексты аналитиков. Колонки: «Характерные признаки» (итоговый), «авто», «правка»:
    нет | да | состав изменился (правка сохранена для другого состава — показывается, но с предупреждением)."""
    saved = load_descriptions().get(key, {})
    t = table.copy()
    t["авто"] = t["Характерные признаки"]
    t["правка"] = "нет"
    for i, r in t.iterrows():
        s = saved.get(r["Кластер"])
        if s and s.get("text"):
            t.at[i, "Характерные признаки"] = s["text"]
            t.at[i, "правка"] = "да" if s.get("состав") == r["состав"] else "состав изменился"
    return t


# ============================================================================ выгрузка
def z_fill(z: float) -> str:
    """Расходящаяся шкала для z: синий (−2) — светло-серый (0) — красный (+2)."""
    if not np.isfinite(z):
        return "FFFFFF"
    t = max(-1.0, min(1.0, z / 2))
    lo, mid, hi = (0x18, 0x4F, 0x95), (0xF0, 0xEF, 0xEC), (0xE3, 0x49, 0x48)
    a, b = (mid, hi) if t >= 0 else (mid, lo)
    rgb = [round(a[i] + (b[i] - a[i]) * abs(t)) for i in range(3)]
    return "".join(f"{v:02X}" for v in rgb)


def table1_excel(res: dict, table: pd.DataFrame, title: str) -> bytes:
    """Excel: лист «Таблица 1» (кластер окрашен своим цветом) и лист «Матрица» (средние, цвет — по z)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Таблица 1"
    ws.append([title])
    ws["A1"].font = Font(bold=True)
    cols = ["Кластер", "Число МО", "Характерные признаки", "Примеры МО"]
    ws.append(cols)
    for c in ws[2]:
        c.font = Font(bold=True)
    for _, r in table.iterrows():
        ws.append([r[c] for c in cols])
        cell = ws.cell(row=ws.max_row, column=1)
        cell.fill = PatternFill("solid", fgColor=cluster_color(int(r["label"]))[1:].upper())
        cell.font = Font(bold=True, color="FFFFFF")
        for col in (3, 4):
            ws.cell(row=ws.max_row, column=col).alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCD", (10, 10, 70, 50)):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("Матрица")
    ws2.append([f"{title}: средние значения признаков (цвет — средний z-score: синий — ниже, красный — выше)"])
    feats = res["features"]
    ws2.append(["Кластер", *[res["names"][f] for f in feats]])
    for c in ws2[2]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True)
    for lab in res["raw"].index:
        ws2.append([clustering.code(lab), *[float(res["raw"].at[lab, f]) for f in feats]])
        ws2.cell(row=ws2.max_row, column=1).fill = PatternFill("solid", fgColor=cluster_color(int(lab))[1:].upper())
        for j, f in enumerate(feats, start=2):
            cell = ws2.cell(row=ws2.max_row, column=j)
            cell.fill = PatternFill("solid", fgColor=z_fill(float(res["z"].at[lab, f])))
            cell.number_format = "#,##0.00"
    ws2.column_dimensions["A"].width = 10
    for j in range(2, len(feats) + 2):
        ws2.column_dimensions[ws2.cell(row=2, column=j).column_letter].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ============================================================================ таблица 2
TRAJECTORIES = ["стабильный", "рост", "снижение", "колебание", "нет данных"]


def default_periods(y0: int, y1: int, n: int = 3) -> list[int]:
    """n равноудалённых лет окна [y0, y1] (по умолчанию — начало, середина, конец)."""
    if n <= 1 or y1 <= y0:
        return [y1]
    return sorted({int(round(y0 + (y1 - y0) * i / (n - 1))) for i in range(n)})


def trajectory(seq) -> str:
    """Траектория по последовательности меток (0 → K1; NaN — нет метки в периоде, пропускается).

    стабильный — одна метка; рост — номер только уменьшается (переход к кластеру с более высоким ВМП);
    снижение — только увеличивается; колебание — и то и другое; нет данных — меток меньше двух.
    """
    s = [int(v) for v in seq if v is not None and not (isinstance(v, float) and np.isnan(v))]
    if len(s) < 2:
        return "нет данных"
    d = np.diff(s)
    if (d == 0).all():
        return "стабильный"
    if (d <= 0).all():
        return "рост"
    if (d >= 0).all():
        return "снижение"
    return "колебание"


def periods_table(labels: pd.DataFrame, periods: list[int]) -> pd.DataFrame:
    """Уровень МО: № | Субъект | ФО | МО | тип МО | метки по периодам (0 → K1, NaN — нет) | траектория.

    Строки — МО, у которых есть метка хотя бы в одном из периодов; порядок — ФО, субъект, МО.
    """
    w = labels[labels["year"].isin(periods)].pivot_table(index="territory_id", columns="year", values="label")
    w = w.reindex(columns=periods)
    mo = pd.read_parquet(PROCESSED / "mo.parquet").set_index("territory_id")
    t = pd.DataFrame(
        {
            "territory_id": w.index,
            "Субъект": w.index.map(mo["region_name"]),
            "ФО": w.index.map(mo["federal_district"]),
            "МО": w.index.map(mo["name"]),
            "тип МО": w.index.map(mo["type"]),
        }
    )
    for y in periods:
        t[y] = w[y].to_numpy()
    t["траектория"] = [trajectory(r) for r in w.to_numpy()]
    t = t.sort_values(["ФО", "Субъект", "МО"]).reset_index(drop=True)
    t.insert(0, "№", np.arange(1, len(t) + 1))
    return t


def subject_table(labels: pd.DataFrame, periods: list[int], by: str = "count") -> pd.DataFrame:
    """Уровень субъектов: доминирующий кластер МО субъекта в каждом периоде и доля в нём.

    by = count — доля МО субъекта; pop — доля населения субъекта (по МО с меткой) в МО доминирующего кластера.
    Колонки: №, Субъект, ФО, МО (число МО с меткой в последнем периоде), <год> (метка), «<год> доля», траектория.
    """
    mo = pd.read_parquet(PROCESSED / "mo.parquet").set_index("territory_id")
    d = labels[labels["year"].isin(periods)].copy()
    d["Субъект"] = d["territory_id"].map(mo["region_name"])
    d["ФО"] = d["territory_id"].map(mo["federal_district"])
    if by == "pop":
        pop = pd.read_parquet(PROCESSED / "indicators_wide.parquet", columns=["territory_id", "year", "pop"])
        d = d.merge(pop, on=["territory_id", "year"], how="left")
        d["w"] = d["pop"].fillna(0.0)
    else:
        d["w"] = 1.0
    g = d.groupby(["Субъект", "year", "label"])["w"].sum().rename("w").reset_index()
    g["share"] = g["w"] / g.groupby(["Субъект", "year"])["w"].transform("sum")
    # доминирующий кластер: наибольшая доля; при равенстве — меньший номер (более высокий ВМП)
    dom = g.sort_values(["Субъект", "year", "share", "label"], ascending=[True, True, False, True])
    dom = dom.drop_duplicates(["Субъект", "year"]).set_index(["Субъект", "year"])
    fd = d.drop_duplicates("Субъект").set_index("Субъект")["ФО"]
    subj = sorted(d["Субъект"].dropna().unique(), key=lambda s: (str(fd[s]), s))
    out = pd.DataFrame({"Субъект": subj, "ФО": [fd[s] for s in subj]})
    last = d[d["year"].eq(max(periods))].groupby("Субъект").size()
    out["МО"] = [int(last.get(s, 0)) for s in subj]
    for y in periods:
        out[y] = [dom["label"].get((s, y), np.nan) for s in subj]
        out[f"{y} доля"] = [dom["share"].get((s, y), np.nan) for s in subj]
    out["траектория"] = [trajectory([out.at[i, y] for y in periods]) for i in out.index]
    out.insert(0, "№", np.arange(1, len(out) + 1))
    return out


def transitions_summary(mo_table: pd.DataFrame) -> pd.DataFrame:
    """Число и доля МО по траекториям: вся выборка («Россия») и федеральные округа."""
    rows = []
    for name, g in [("Россия", mo_table), *sorted(mo_table.groupby("ФО"), key=lambda x: str(x[0]))]:
        vc = g["траектория"].value_counts()
        n = len(g)
        r = {"территория": name, "МО": n}
        for tr in TRAJECTORIES:
            r[tr] = int(vc.get(tr, 0))
            r[f"{tr}, %"] = 100 * vc.get(tr, 0) / n if n else np.nan
        rows.append(r)
    return pd.DataFrame(rows)


def label_text(v) -> str:
    """Метка ячейки: K1…Kn или «—» (нет метки)."""
    return "—" if v is None or (isinstance(v, float) and np.isnan(v)) else clustering.code(int(v))


def tint(hex_color: str, a: float = 0.35) -> str:
    """Светлый оттенок цвета для фона строки (смешение с белым), RRGGBB."""
    h = hex_color.lstrip("#")
    rgb = [int(h[i : i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(255 - (255 - c) * a):02X}" for c in rgb)


def table2_excel(
    mo_table: pd.DataFrame,
    subj_count: pd.DataFrame,
    subj_pop: pd.DataFrame,
    summ: pd.DataFrame,
    periods: list[int],
    title: str,
) -> bytes:
    """Excel: «МО» (метки — цвет кластера, строки — оттенок цвета траектории), «Субъекты (по числу МО)»,
    «Субъекты (по населению)», «Переходы»."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    tcol = cfg()["periods"]["trajectory_colors"]
    wb = Workbook()

    def header(ws, cols):
        ws.append([title])
        ws["A1"].font = Font(bold=True)
        ws.append([str(c) for c in cols])
        for c in ws[2]:
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "A3"

    def label_cell(cell, v):
        cell.value = label_text(v)
        cell.alignment = Alignment(horizontal="center")
        if cell.value != "—":
            cell.fill = PatternFill("solid", fgColor=cluster_color(int(v))[1:].upper())
            cell.font = Font(bold=True, color="FFFFFF")

    ws = wb.active
    ws.title = "МО"
    cols = ["№", "Субъект", "МО", "тип МО", *periods, "траектория"]
    header(ws, cols)
    for _, r in mo_table.iterrows():
        ws.append([r["№"], r["Субъект"], r["МО"], r["тип МО"], *[None] * len(periods), r["траектория"]])
        i = ws.max_row
        row_fill = PatternFill("solid", fgColor=tint(tcol.get(r["траектория"], "#ffffff")))
        for j in range(1, len(cols) + 1):
            ws.cell(row=i, column=j).fill = row_fill
        for j, y in enumerate(periods, start=5):
            label_cell(ws.cell(row=i, column=j), r[y])
    for col, wdt in zip("ABCD", (6, 30, 40, 22)):
        ws.column_dimensions[col].width = wdt
    ws.auto_filter.ref = f"A2:{ws.cell(row=2, column=len(cols)).column_letter}{ws.max_row}"

    for name, st in (("Субъекты (по числу МО)", subj_count), ("Субъекты (по населению)", subj_pop)):
        ws = wb.create_sheet(name)
        cols = ["№", "Субъект", "ФО", "МО", *[c for y in periods for c in (y, f"{y} доля")], "траектория"]
        header(ws, cols)
        for _, r in st.iterrows():
            ws.append([r["№"], r["Субъект"], r["ФО"], r["МО"], *[None] * (2 * len(periods)), r["траектория"]])
            i = ws.max_row
            for j, y in enumerate(periods):
                label_cell(ws.cell(row=i, column=5 + 2 * j), r[y])
                sh = r[f"{y} доля"]
                c = ws.cell(row=i, column=6 + 2 * j, value=None if pd.isna(sh) else float(sh))
                c.number_format = "0%"
            ws.cell(row=i, column=len(cols)).fill = PatternFill(
                "solid", fgColor=tint(tcol.get(r["траектория"], "#ffffff"))
            )
        ws.column_dimensions["B"].width = 34

    ws = wb.create_sheet("Переходы")
    header(ws, list(summ.columns))
    for _, r in summ.iterrows():
        ws.append([None if isinstance(v, float) and np.isnan(v) else v for v in r.tolist()])
        for j, c in enumerate(summ.columns, start=1):
            if str(c).endswith("%"):
                ws.cell(row=ws.max_row, column=j).number_format = "0.0"
    ws.column_dimensions["A"].width = 12
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def table2_html(subj: pd.DataFrame, summ: pd.DataFrame, periods: list[int], title: str, note: str = "") -> str:
    """Страница для печати (или «Сохранить как PDF» в браузере): субъекты с цветами и сводка переходов."""
    import html

    tcol = cfg()["periods"]["trajectory_colors"]

    def cell(v, share):
        t = label_text(v)
        if t == "—":
            return '<td class="c">—</td>'
        sh = "" if pd.isna(share) else f'<span class="s">{share:.0%}</span>'
        return f'<td class="c" style="background:{cluster_color(int(v))};color:#fff"><b>{t}</b> {sh}</td>'

    rows = []
    for _, r in subj.iterrows():
        tr = r["траектория"]
        rows.append(
            f"<tr><td>{r['№']}</td><td>{html.escape(str(r['Субъект']))}</td><td>{html.escape(str(r['ФО']))}</td>"
            f"<td class='n'>{r['МО']}</td>"
            + "".join(cell(r[y], r[f"{y} доля"]) for y in periods)
            + f"<td style='background:#{tint(tcol.get(tr, '#ffffff'))}'>{tr}</td></tr>"
        )
    srows = "".join(
        "<tr>"
        + f"<td>{html.escape(str(r['территория']))}</td><td class='n'>{r['МО']}</td>"
        + "".join(f"<td class='n'>{r[t]} ({r[t + ', %']:.0f}%)</td>" for t in TRAJECTORIES)
        + "</tr>"
        for _, r in summ.iterrows()
    )
    head = "".join(f"<th>{y}</th>" for y in periods)
    thead = "".join(f"<th>{t}</th>" for t in TRAJECTORIES)
    css = (
        "body{font:12px/1.35 Arial,sans-serif;color:#222;margin:16px} h1{font-size:16px} h2{font-size:14px;margin-top:20px}"
        "table{border-collapse:collapse;width:100%} th,td{border:1px solid #ccc;padding:3px 5px;vertical-align:top}"
        "th{background:#f0efec} td.c{text-align:center;white-space:nowrap} td.n{text-align:right} .s{font-size:10px}"
        "p.note{color:#555} @media print{body{margin:0} tr{page-break-inside:avoid}}"
    )
    return (
        f'<!doctype html><html lang="ru"><head><meta charset="utf-8"><title>{html.escape(title)}</title>'
        f"<style>{css}</style></head><body><h1>{html.escape(title)}</h1><p class='note'>{html.escape(note)}</p>"
        f"<table><thead><tr><th>№</th><th>Субъект</th><th>ФО</th><th>МО</th>{head}<th>траектория</th></tr></thead>"
        f"<tbody>{''.join(rows)}</tbody></table>"
        "<h2>Сводка переходов (число МО и доля по траекториям)</h2>"
        f"<table><thead><tr><th>территория</th><th>МО</th>{thead}</tr></thead><tbody>{srows}</tbody></table>"
        "</body></html>"
    )


def row_css(traj: str) -> str:
    """CSS строки таблицы 2 в интерфейсе: светлый оттенок цвета траектории и тёмный текст (читается и в тёмной
    теме Streamlit); «стабильный» и «нет данных» — без фона и цвета, как обычные строки темы."""
    if traj in ("стабильный", "нет данных"):
        return ""
    c = cfg()["periods"]["trajectory_colors"].get(traj)
    return f"background-color: #{tint(c)}; color: #1f1f1f" if c else ""
