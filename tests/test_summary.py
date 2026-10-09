"""Сводная таблица 1 «Характерные признаки кластеров»."""

from __future__ import annotations

import io

import numpy as np
import pandas as pd
import pytest

from src import network, summary
from src.io import PROCESSED

needs_ind = pytest.mark.skipif(not (PROCESSED / "indicators_wide.parquet").exists(), reason="нет показателей")
CFG = {"thresholds": {"extreme": 1.0, "high": 0.5}, "n_features": [3, 5], "n_examples": 4}


@pytest.fixture
def zm():
    # кластеры × признаки: средние z
    return pd.DataFrame(
        {"a": [1.5, 1.2, -0.2], "b": [-1.3, 0.1, -1.1], "c": [0.6, -0.7, 0.3], "d": [0.1, 0.05, 0.2], "e": [0.2, 0, 0]},
        index=[0, 1, 2],
    )


def test_category_rules(zm):
    th = CFG["thresholds"]
    assert summary.category(zm, 0, "a", th) == "max"  # наибольшее среднее и z ≥ 1
    assert summary.category(zm, 1, "a", th) == "high"  # z ≥ 1, но максимум у другого кластера → «высокие»
    assert summary.category(zm, 0, "b", th) == "min"
    assert summary.category(zm, 2, "b", th) == "low"  # z ≤ −1, но минимум у кластера 0
    assert summary.category(zm, 1, "c", th) == "low"
    assert summary.category(zm, 0, "c", th) == "high"
    assert summary.category(zm, 2, "a", th) == "mid"


def test_describe_picks_3_to_5_by_abs_z(zm):
    names = {c: c.upper() for c in zm.columns}
    text, sel = summary.describe(zm, 0, names, CFG)
    assert sel == ["a", "b", "c"]  # три признака с |z| ≥ 0,5, по убыванию |z|
    assert text == "Максимальные значения: A. Высокие значения: C. Минимальные значения: B."
    text2, sel2 = summary.describe(zm, 2, names, CFG)
    assert len(sel2) == 3 and sel2[0] == "b"  # сильный один, добираем до трёх
    assert "Близкие к среднему" in text2


def test_short_name():
    assert summary.short_name("ВМП на душу населения (оценка команды)") == "ВМП на душу населения"
    assert summary.short_name("Плотность населения") == "плотность населения"


def test_saved_text_kept_and_warns_when_members_change(tmp_path, monkeypatch):
    monkeypatch.setattr(summary, "_desc_path", lambda: tmp_path / "desc.yaml")
    t = pd.DataFrame({"Кластер": ["K1", "K2"], "Характерные признаки": ["авто 1", "авто 2"], "состав": ["aaa", "bbb"]})
    summary.save_description("key", "K1", "ручной текст", "aaa", "авто 1")
    summary.save_description("key", "K2", "старый текст", "zzz", "авто 2")
    out = summary.apply_descriptions(t, "key")
    assert list(out["Характерные признаки"]) == ["ручной текст", "старый текст"]
    assert list(out["правка"]) == ["да", "состав изменился"]
    assert list(out["авто"]) == ["авто 1", "авто 2"]
    assert summary.apply_descriptions(t, "другое разбиение")["правка"].eq("нет").all()


def test_fingerprint_order_independent():
    assert summary.fingerprint([3, 1, 2]) == summary.fingerprint([1, 2, 3]) != summary.fingerprint([1, 2])


@needs_ind
def test_table1_on_dfo_and_excel_colors():
    from openpyxl import load_workbook

    p = network.default_params()
    p["sample"]["federal_districts"] = ["ДФО"]
    lab = summary.partition(p, "ward_kmeans", 4, "pooled")
    res = summary.characteristic_table(lab, p)
    t = res["table"]
    assert list(t["Кластер"]) == ["K1", "K2", "K3", "K4"]
    assert t["Число МО"].sum() == lab["year"].eq(res["year"]).sum()
    assert t["Характерные признаки"].str.len().gt(0).all()
    # K1 — самый высокий ВМП на душу (медиана; средний z тоже наибольший)
    assert res["z"]["gmp_pc"].idxmax() == 0
    assert np.isfinite(res["raw"].to_numpy()).any()
    wb = load_workbook(io.BytesIO(summary.table1_excel(res, t, "тест")))
    assert wb["Таблица 1"]["A3"].fill.fgColor.rgb.endswith(summary.CLUSTER_COLORS[0][1:].upper())
    assert wb["Матрица"].max_row == 2 + len(t)


# ---------------------------------------------------------------- таблица 2
@pytest.mark.parametrize(
    "seq, expected",
    [
        ([2, 2, 2], "стабильный"),
        ([3, 2, 0], "рост"),  # номер уменьшается → кластер с более высоким ВМП
        ([1, 1, 4], "снижение"),
        ([1, 3, 1], "колебание"),
        ([2, np.nan, 1], "рост"),  # период без метки пропускается
        ([np.nan, 2, np.nan], "нет данных"),
    ],
)
def test_trajectory(seq, expected):
    assert summary.trajectory(seq) == expected


def test_default_periods():
    assert summary.default_periods(2017, 2024) == [2017, 2020, 2024]
    assert summary.default_periods(2013, 2023) == [2013, 2018, 2023]
    assert summary.default_periods(2017, 2024, 4) == [2017, 2019, 2022, 2024]


@needs_ind
def test_table2_on_dfo():
    from openpyxl import load_workbook

    p = network.default_params()
    p["sample"]["federal_districts"] = ["ДФО"]
    lab = summary.partition(p, "ward_kmeans", 4, "pooled")
    per = [2017, 2020, 2024]
    mt = summary.periods_table(lab, per)
    assert list(mt.columns) == ["№", "territory_id", "Субъект", "ФО", "МО", "тип МО", *per, "траектория"]
    assert mt["territory_id"].is_unique and set(mt["траектория"]) <= set(summary.TRAJECTORIES)
    # траектория пересчитывается из меток строки
    r = mt.iloc[0]
    assert r["траектория"] == summary.trajectory([r[y] for y in per])
    for by in ("count", "pop"):
        st = summary.subject_table(lab, per, by)
        assert st["Субъект"].nunique() == mt["Субъект"].nunique() == 11  # субъекты ДФО
        for y in per:
            sh = st[f"{y} доля"].dropna()
            assert sh.between(1 / 4 - 1e-9, 1).all()  # доминирующий кластер — не меньше 1/k
    # доля по числу МО — ручной пересчёт для одного субъекта
    st = summary.subject_table(lab, per, "count")
    s0, y = st.iloc[0], per[-1]
    sub = mt[mt["Субъект"].eq(s0["Субъект"])][y].dropna()
    assert s0[y] == sub.value_counts().sort_index().idxmax()
    assert s0[f"{y} доля"] == pytest.approx(sub.value_counts().max() / len(sub))
    sm = summary.transitions_summary(mt)
    assert sm.iloc[0]["МО"] == len(mt) and sm.iloc[0][summary.TRAJECTORIES].sum() == len(mt)
    wb = load_workbook(io.BytesIO(summary.table2_excel(mt, st, st, sm, per, "тест")))
    assert wb.sheetnames == ["МО", "Субъекты (по числу МО)", "Субъекты (по населению)", "Переходы"]
    ws = wb["МО"]
    lab_cell = ws.cell(row=3, column=5)
    assert lab_cell.value.startswith("K") or lab_cell.value == "—"
    if lab_cell.value != "—":
        assert lab_cell.fill.fgColor.rgb.endswith(summary.cluster_color(int(lab_cell.value[1:]) - 1)[1:].upper())
    assert ws.cell(row=3, column=2).fill.fgColor.rgb not in ("00000000", None)  # строка окрашена по траектории
    html = summary.table2_html(st, sm, per, "тест")
    assert html.count("<tr>") >= len(st) and "K" in html


def test_row_css_readable_in_dark_theme():
    """Строки «стабильный» и «нет данных» — без фона (наследуют тему); цветные — с явным тёмным текстом."""
    assert summary.row_css("стабильный") == "" and summary.row_css("нет данных") == ""
    for tr in ("рост", "снижение", "колебание"):
        assert "color: #1f1f1f" in summary.row_css(tr) and "background-color" in summary.row_css(tr)
